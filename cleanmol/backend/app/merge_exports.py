"""
Merge JSONL exports into unified dataset files.

Reads all JSONL files from output/exports/ and produces:
1. unified_dataset.xlsx - Excel workbook with 5 sheets
2. analysis_ready.csv - Flat denormalized format for ML
3. molecules.smi - SMILES file for Quat Generator
4. fairchem_uma_candidates.csv - Review sheet for UMA/OMol atomistic modeling
5. candidate_generation_seed.smi - Modern disinfectant seed molecules for generation
6. legacy_or_low_priority_molecules.csv - Molecules kept out of the seed set
7. screening_score_profile.json - Multi-objective discovery scoring profile
"""
import json
import csv
import re
from pathlib import Path
from typing import Dict, Any, List, Optional, Tuple
from collections import defaultdict
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from .discovery_filters import (
    enrich_modern_discovery_fields,
    is_generation_seed,
    is_legacy_or_low_priority,
)


# Style definitions (matching excel_export.py)
HEADER_FONT = Font(bold=True, color="FFFFFF")
HEADER_FILL = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
HEADER_ALIGNMENT = Alignment(horizontal="center", vertical="center", wrap_text=True)
ALT_ROW_FILL = PatternFill(start_color="D9E2F3", end_color="D9E2F3", fill_type="solid")
THIN_BORDER = Border(
    left=Side(style="thin"),
    right=Side(style="thin"),
    top=Side(style="thin"),
    bottom=Side(style="thin")
)
NO_SMILES_FILL = PatternFill(start_color="FFEB9C", end_color="FFEB9C", fill_type="solid")


def _auto_width(ws, min_width: int = 10, max_width: int = 50):
    """Auto-fit column widths based on content."""
    for col_idx, column_cells in enumerate(ws.columns, 1):
        max_len = 0
        for cell in column_cells:
            try:
                cell_len = len(str(cell.value or ""))
                if cell_len > max_len:
                    max_len = cell_len
            except Exception:
                pass
        adjusted_width = min(max(max_len + 2, min_width), max_width)
        ws.column_dimensions[get_column_letter(col_idx)].width = adjusted_width


def _apply_header_style(ws, row: int = 1):
    """Apply header styling to specified row."""
    for cell in ws[row]:
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = HEADER_ALIGNMENT
        cell.border = THIN_BORDER


def _apply_data_style(ws, start_row: int = 2):
    """Apply data row styling with alternating colors."""
    for row_idx, row in enumerate(ws.iter_rows(min_row=start_row), start=0):
        for cell in row:
            cell.border = THIN_BORDER
            cell.alignment = Alignment(vertical="center", wrap_text=True)
            if row_idx % 2 == 1:
                cell.fill = ALT_ROW_FILL


def _read_jsonl(path: Path) -> List[Dict[str, Any]]:
    """Read a JSONL file and return list of records."""
    records = []
    if not path.exists():
        return records
    with path.open("r", encoding="utf-8") as f:
        for line in f:
            line = line.strip()
            if line:
                try:
                    records.append(json.loads(line))
                except json.JSONDecodeError:
                    pass
    return records


def _get_source_paper(output_dir: Path, doc_id: str) -> str:
    """Get source paper filename from bundle metadata."""
    metadata_path = output_dir / "bundles" / doc_id / "metadata.json"
    if metadata_path.exists():
        try:
            metadata = json.loads(metadata_path.read_text(encoding="utf-8"))
            filename = metadata.get("filename", "")
            # Strip .pdf extension
            if filename.lower().endswith(".pdf"):
                filename = filename[:-4]
            return filename
        except (json.JSONDecodeError, IOError):
            pass
    return doc_id  # Fallback to doc_id if metadata not found


def _infer_missing_smiles_reason(molecule: Dict[str, Any]) -> str:
    """Infer reason why SMILES is missing based on molecule name."""
    name = molecule.get("name_as_written", "") or molecule.get("normalized_name", "") or ""
    name_lower = name.lower()

    if "polymyxin" in name_lower or "peptide" in name_lower:
        return "Complex peptide"
    if "triton" in name_lower or "mixture" in name_lower:
        return "Polymer/mixture"
    return "Not found in chemical databases"


def _normalize_organism(organism: str) -> str:
    """Normalize organism names for column matching."""
    if not organism:
        return ""
    org_lower = organism.lower()

    if "aureus" in org_lower or "s. aureus" in org_lower or "staphylococcus" in org_lower:
        return "S_aureus"
    if "coli" in org_lower or "e. coli" in org_lower or "escherichia" in org_lower:
        return "E_coli"
    if "aeruginosa" in org_lower or "p. aeruginosa" in org_lower or "pseudomonas" in org_lower:
        return "P_aeruginosa"
    if "albicans" in org_lower or "c. albicans" in org_lower or "candida" in org_lower:
        return "C_albicans"

    return ""


def _collect_exports(output_dir: Path) -> Tuple[
    Dict[str, List[Dict]],  # molecules by doc_id
    Dict[str, List[Dict]],  # experiments by doc_id
    Dict[str, List[Dict]],  # results by doc_id
    List[str]               # doc_ids found
]:
    """Collect all JSONL exports grouped by doc_id."""
    exports_dir = output_dir / "exports"
    if not exports_dir.exists():
        return {}, {}, {}, []

    # Group files by doc_id
    molecules_by_doc = defaultdict(list)
    experiments_by_doc = defaultdict(list)
    results_by_doc = defaultdict(list)
    doc_ids = set()

    # Pattern: {doc_id}_{timestamp}_{type}.jsonl
    for jsonl_file in exports_dir.glob("*.jsonl"):
        filename = jsonl_file.stem  # Without .jsonl
        parts = filename.rsplit("_", 2)  # Split from right to get doc_id_timestamp_type

        if len(parts) < 3:
            continue

        file_type = parts[-1]  # molecules, experiments, results, dataset
        if file_type not in ("molecules", "experiments", "results"):
            continue

        # Reconstruct doc_id (everything before _timestamp_type)
        # Format: doc_id_YYYYMMDD_HHMMSS_type
        # So we need to find the doc_id part
        # Split by underscore and find timestamp pattern
        all_parts = filename.split("_")
        timestamp_idx = None
        for i, part in enumerate(all_parts):
            if len(part) == 8 and part.isdigit():  # YYYYMMDD
                timestamp_idx = i
                break

        if timestamp_idx is None:
            continue

        doc_id = "_".join(all_parts[:timestamp_idx])
        doc_ids.add(doc_id)

        records = _read_jsonl(jsonl_file)
        if file_type == "molecules":
            molecules_by_doc[doc_id].extend(records)
        elif file_type == "experiments":
            experiments_by_doc[doc_id].extend(records)
        elif file_type == "results":
            results_by_doc[doc_id].extend(records)

    return molecules_by_doc, experiments_by_doc, results_by_doc, sorted(doc_ids)


def _serialize_field(value: Any) -> str:
    """Serialize a field value for Excel/CSV output."""
    if value is None:
        return ""
    if isinstance(value, (list, dict)):
        return json.dumps(value, ensure_ascii=False)
    return str(value)


def _sanitize_cell(value: Any) -> Any:
    """Sanitize a single cell value for Excel compatibility.

    Converts dict/list values to JSON strings to prevent Excel errors.
    """
    if value is None:
        return ""
    if isinstance(value, (dict, list)):
        return json.dumps(value, ensure_ascii=False)
    return value


def _sanitize_row(row: List[Any]) -> List[Any]:
    """Sanitize an entire row for Excel compatibility."""
    return [_sanitize_cell(cell) for cell in row]


def _infer_uma_readiness(molecule: Dict[str, Any]) -> Tuple[str, str]:
    """Classify whether a molecule is ready to become a FairChem UMA input."""
    smiles = molecule.get("smiles")
    if not smiles:
        return "missing_smiles", "No SMILES available."

    if molecule.get("smiles_valid") is False:
        return "invalid_smiles", molecule.get("smiles_error") or "SMILES failed validation."

    fragment_count = molecule.get("fragment_count")
    formal_charge = molecule.get("formal_charge")

    if fragment_count is None or formal_charge is None:
        return (
            "needs_structure_review",
            "RDKit charge/fragment metadata is missing; rerun validation with RDKit before preparing UMA inputs.",
        )

    if fragment_count is not None:
        try:
            if int(fragment_count) > 1:
                return (
                    "needs_fragment_review",
                    "Disconnected fragments or salt/counterion detected; decide whether UMA should model the full explicit salt or a selected active fragment.",
                )
        except (TypeError, ValueError):
            pass

    if formal_charge is not None:
        try:
            if int(formal_charge) != 0:
                return (
                    "needs_charge_review",
                    "Charged molecule; confirm total_charge and spin_multiplicity before using UMA task omol.",
                )
        except (TypeError, ValueError):
            pass

    return "candidate", "Candidate for RDKit 3D conformer generation and UMA task omol; confirm spin multiplicity."


def _build_screening_score_profile(stats: Dict[str, Any]) -> Dict[str, Any]:
    """Create a model-agnostic score profile for Chemprop/REINVENT candidate ranking."""
    return {
        "profile_name": "lysol_2_modern_disinfectant_discovery",
        "version": "2026-05-04",
        "purpose": (
            "Prioritize modern cationic/amphiphilic disinfectant candidates for chemist review, "
            "not to make claims of safety, efficacy, or synthesizability."
        ),
        "recommended_stack": {
            "activity_predictor": "Chemprop v2 multitask ensemble with a Morgan-fingerprint baseline",
            "generator": "REINVENT 4 for de novo design, scaffold hopping, R-group replacement, and multi-objective optimization",
            "physics_filter": "FAIRChem UMA with task omol, using OMol25-style charge/spin-aware molecular inputs",
            "primary_activity_data": "CleanMol-curated QAC/biocide literature plus PubChem BioAssay and ChEMBL antimicrobial/cytotoxicity rows",
            "toxicity_selectivity_data": "Tox21/ToxCast, EPA CompTox, ChEMBL cytotoxicity, and hemolysis/selectivity assays",
        },
        "seed_policy": {
            "include": [
                "candidate_tier == modern_seed",
                "generation_recommendation == use_as_generation_seed",
                "valid or RDKit-skipped SMILES with a stable cationic/amphiphilic head group",
            ],
            "review_before_include": [
                "candidate_tier == needs_review",
                "charged salts with disconnected fragments",
                "missing charge, counterion, or formulation context",
            ],
            "exclude_as_primary_seeds": [
                "neutral single-nitrogen compounds",
                "simple amines without a stable cationic head group",
                "invalid or missing SMILES",
                "molecules with no modern scaffold tags unless a chemist explicitly promotes them",
            ],
        },
        "reward_terms": [
            {"name": "broad_spectrum_predicted_activity", "direction": "maximize", "source": "Chemprop multitask output"},
            {"name": "selectivity_margin", "direction": "maximize", "source": "toxicity/hemolysis vs antimicrobial potency"},
            {"name": "modern_cationic_amphiphilic_scaffold", "direction": "maximize", "source": "CleanMol modern_disinfectant_score"},
            {"name": "gemini_or_bis_cationic_diversity", "direction": "reward_moderately", "source": "scaffold tags"},
            {"name": "novelty_near_known_actives", "direction": "balance", "source": "nearest-neighbor similarity"},
            {"name": "formulation_compatible_property_window", "direction": "maximize", "source": "RDKit descriptors and reviewer rules"},
        ],
        "penalty_terms": [
            {"name": "neutral_single_nitrogen_or_simple_amine", "direction": "strong_penalty"},
            {"name": "reactive_or_structural_liability_alert", "direction": "strong_penalty"},
            {"name": "predicted_cytotoxicity_or_hemolysis", "direction": "strong_penalty"},
            {"name": "extreme_size_charge_or_lipophilicity", "direction": "penalty"},
            {"name": "uma_charge_fragment_uncertainty", "direction": "review_or_penalty"},
        ],
        "required_lab_packet_fields": [
            "SMILES and structure image",
            "candidate_tier and generation_recommendation",
            "nearest known analogs",
            "activity prediction with uncertainty",
            "toxicity/selectivity prediction with uncertainty",
            "literature evidence trail",
            "UMA readiness, total charge, spin multiplicity, and fragment notes",
            "chemist review notes before any lab work",
        ],
        "current_export_stats": stats,
    }


def merge_and_export_datasets(output_dir: Path, logger=None) -> dict:
    """
    Merge all JSONL exports into unified dataset files.

    Args:
        output_dir: Path to output directory containing exports/ and bundles/
        logger: Optional PipelineLogger for status updates

    Returns:
        dict with paths to created files and statistics
    """
    output_dir = Path(output_dir)

    if logger:
        logger.emit_progress("MERGE", "start")

    # Collect all exports
    molecules_by_doc, experiments_by_doc, results_by_doc, doc_ids = _collect_exports(output_dir)

    if not doc_ids:
        if logger:
            logger.log_warning("No JSONL exports found to merge", stage="Merge Exports")
        return {"status": "no_data", "files_created": []}

    # Build source paper mapping
    source_papers = {}
    for doc_id in doc_ids:
        source_papers[doc_id] = _get_source_paper(output_dir, doc_id)

    # Aggregate all data with source paper info
    all_molecules = []
    all_experiments = []
    all_results = []

    for doc_id in doc_ids:
        source_paper = source_papers[doc_id]

        for mol in molecules_by_doc.get(doc_id, []):
            mol["source_paper"] = source_paper
            mol["doc_id"] = doc_id
            all_molecules.append(mol)

        for exp in experiments_by_doc.get(doc_id, []):
            exp["source_paper"] = source_paper
            exp["doc_id"] = doc_id
            all_experiments.append(exp)

        for res in results_by_doc.get(doc_id, []):
            res["source_paper"] = source_paper
            res["doc_id"] = doc_id
            all_results.append(res)

    # Attach modern discovery profile fields before any export is written.
    # This keeps old/simple nitrogen compounds out of the generation seed file
    # while preserving them as baselines or activity references.
    for mol in all_molecules:
        enrich_modern_discovery_fields(mol)

    # Build experiment lookup for organism info
    experiment_lookup = {}
    for exp in all_experiments:
        exp_id = exp.get("experiment_id")
        doc_id = exp.get("doc_id")
        if exp_id and doc_id:
            key = f"{doc_id}:{exp_id}"
            experiment_lookup[key] = exp

    # Statistics
    molecules_with_smiles = [m for m in all_molecules if m.get("smiles")]
    molecules_without_smiles = [m for m in all_molecules if not m.get("smiles")]

    files_created = []
    errors = []

    # Compute stats early for use in all outputs
    # Deduplicate by SMILES for counting
    seen_smiles_for_count = set()
    for mol in all_molecules:
        smiles = mol.get("smiles", "")
        if smiles:
            seen_smiles_for_count.add(smiles)

    stats = {
        "molecules_with_smiles": len(molecules_with_smiles),
        "molecules_without_smiles": len(molecules_without_smiles),
        "unique_smiles": len(seen_smiles_for_count),
        "modern_seed_candidates": sum(1 for m in all_molecules if is_generation_seed(m)),
        "legacy_or_low_priority": sum(1 for m in all_molecules if m.get("candidate_tier") == "legacy_or_low_priority"),
        "experiments": len(all_experiments),
        "results": len(all_results),
        "papers_processed": len(doc_ids)
    }

    # ========================================
    # 1. Create unified_dataset.xlsx
    # ========================================
    excel_path = output_dir / "unified_dataset.xlsx"
    try:
        wb = Workbook()
        if "Sheet" in wb.sheetnames:
            del wb["Sheet"]

        # --- Sheet 1: Summary ---
        ws_summary = wb.create_sheet("Summary", 0)
        ws_summary.append(["Field", "Value"])
        ws_summary.append(["Total molecules (with SMILES)", len(molecules_with_smiles)])
        ws_summary.append(["Total molecules (without SMILES)", len(molecules_without_smiles)])
        ws_summary.append(["Modern generation seed candidates", stats["modern_seed_candidates"]])
        ws_summary.append(["Legacy/low-priority molecules", stats["legacy_or_low_priority"]])
        ws_summary.append(["Total experiments", len(all_experiments)])
        ws_summary.append(["Total results", len(all_results)])
        ws_summary.append(["Papers processed", ", ".join(source_papers.values())])

        _apply_header_style(ws_summary, row=1)
        _auto_width(ws_summary)

        # --- Sheet 2: Molecules (WITH SMILES only, deduplicated) ---
        ws_mol = wb.create_sheet("Molecules")
        mol_headers = [
            "source_paper", "molecule_id", "name_as_written", "normalized_name",
            "smiles", "smiles_source", "smiles_confidence", "head_group_class",
            "chain_lengths", "scaffold_class", "modern_scaffold_tags",
            "legacy_flags", "modern_disinfectant_score", "candidate_tier",
            "generation_recommendation", "generation_notes", "notes"
        ]
        ws_mol.append(mol_headers)

        # Deduplicate by SMILES - combine source_papers
        smiles_to_molecules = defaultdict(list)
        for mol in molecules_with_smiles:
            smiles = mol.get("smiles", "")
            if smiles:
                smiles_to_molecules[smiles].append(mol)

        for smiles, mols in smiles_to_molecules.items():
            # Combine source papers
            sources = sorted(set(m.get("source_paper", "") for m in mols))
            combined_source = ", ".join(s for s in sources if s)

            # Use first molecule as representative
            mol = mols[0]
            ws_mol.append(_sanitize_row([
                combined_source,
                mol.get("molecule_id", ""),
                mol.get("name_as_written", ""),
                mol.get("normalized_name", ""),
                smiles,
                mol.get("smiles_source", ""),
                mol.get("smiles_confidence", ""),
                mol.get("head_group_class", ""),
                mol.get("chain_lengths"),
                mol.get("scaffold_class", ""),
                mol.get("modern_scaffold_tags"),
                mol.get("legacy_flags"),
                mol.get("modern_disinfectant_score", ""),
                mol.get("candidate_tier", ""),
                mol.get("generation_recommendation", ""),
                mol.get("generation_notes", ""),
                mol.get("notes", "")
            ]))

        _apply_header_style(ws_mol)
        _apply_data_style(ws_mol)
        _auto_width(ws_mol)

        # --- Sheet 3: Experiments ---
        ws_exp = wb.create_sheet("Experiments")
        exp_headers = [
            "source_paper", "experiment_id", "organism", "strain", "assay_type",
            "conditions", "exposure_protocol", "notes"
        ]
        ws_exp.append(exp_headers)

        for exp in all_experiments:
            ws_exp.append(_sanitize_row([
                exp.get("source_paper", ""),
                exp.get("experiment_id", ""),
                exp.get("organism", ""),
                exp.get("strain", ""),
                exp.get("assay_type", "") or exp.get("type", ""),
                exp.get("conditions"),
                exp.get("exposure_protocol", ""),
                exp.get("notes", "")
            ]))

        _apply_header_style(ws_exp)
        _apply_data_style(ws_exp)
        _auto_width(ws_exp)

        # --- Sheet 4: Results ---
        ws_res = wb.create_sheet("Results")
        res_headers = [
            "source_paper", "result_id", "molecule_id", "experiment_id",
            "endpoint", "value", "units", "directionality", "confidence", "notes"
        ]
        ws_res.append(res_headers)

        for res in all_results:
            ws_res.append(_sanitize_row([
                res.get("source_paper", ""),
                res.get("result_id", ""),
                res.get("molecule_id", ""),
                res.get("experiment_id", ""),
                res.get("endpoint", "") or res.get("property_name", ""),
                res.get("value", ""),
                res.get("units", "") or res.get("unit", ""),
                res.get("directionality", ""),
                res.get("confidence", ""),
                res.get("notes", "")
            ]))

        _apply_header_style(ws_res)
        _apply_data_style(ws_res)
        _auto_width(ws_res)

        # --- Sheet 5: Missing SMILES ---
        ws_missing = wb.create_sheet("Missing SMILES")
        missing_headers = [
            "source_paper", "molecule_id", "name_as_written", "normalized_name", "reason"
        ]
        ws_missing.append(missing_headers)

        for mol in molecules_without_smiles:
            reason = _infer_missing_smiles_reason(mol)
            ws_missing.append(_sanitize_row([
                mol.get("source_paper", ""),
                mol.get("molecule_id", ""),
                mol.get("name_as_written", ""),
                mol.get("normalized_name", ""),
                reason
            ]))

        _apply_header_style(ws_missing)
        _apply_data_style(ws_missing)
        _auto_width(ws_missing)

        # Apply yellow highlight to Missing SMILES rows
        for row_idx, row in enumerate(ws_missing.iter_rows(min_row=2), start=2):
            for cell in row:
                cell.fill = NO_SMILES_FILL

        # Save Excel
        wb.save(excel_path)
        files_created.append(str(excel_path))
    except Exception as e:
        error_msg = f"Excel export failed: {e}"
        errors.append(error_msg)
        if logger:
            logger.log_warning(error_msg, stage="Merge Exports")

    # ========================================
    # 2. Create analysis_ready.csv
    # ========================================
    csv_path = output_dir / "analysis_ready.csv"
    try:
        # Build molecule lookup by (doc_id, molecule_id)
        molecule_lookup = {}
        for mol in molecules_with_smiles:
            mol_id = mol.get("molecule_id")
            doc_id = mol.get("doc_id")
            if mol_id and doc_id:
                key = f"{doc_id}:{mol_id}"
                molecule_lookup[key] = mol

        # Collect MIC and HC50 values per molecule
        molecule_data = defaultdict(lambda: {
            "MIC_S_aureus": None,
            "MIC_E_coli": None,
            "MIC_P_aeruginosa": None,
            "MIC_C_albicans": None,
            "HC50": None
        })

        for res in all_results:
            doc_id = res.get("doc_id")
            mol_id = res.get("molecule_id")
            exp_id = res.get("experiment_id")
            endpoint_val = res.get("endpoint") or res.get("property_name") or ""
            # Handle dict endpoints
            if isinstance(endpoint_val, dict):
                endpoint_val = endpoint_val.get("description", "") or str(endpoint_val)
            endpoint = str(endpoint_val).upper()
            value = res.get("value")

            if not mol_id or not doc_id:
                continue

            mol_key = f"{doc_id}:{mol_id}"

            # Check if this molecule has SMILES
            if mol_key not in molecule_lookup:
                continue

            # Handle HC50
            if "HC50" in endpoint or "HEMOLYSIS" in endpoint:
                if value is not None:
                    try:
                        molecule_data[mol_key]["HC50"] = float(value)
                    except (ValueError, TypeError):
                        pass
                continue

            # Handle MIC values
            if "MIC" in endpoint:
                # Get organism from experiment
                exp_key = f"{doc_id}:{exp_id}"
                exp = experiment_lookup.get(exp_key, {})
                organism = exp.get("organism", "")
                normalized_org = _normalize_organism(organism)

                if normalized_org and value is not None:
                    col_name = f"MIC_{normalized_org}"
                    if col_name in molecule_data[mol_key]:
                        try:
                            val = float(value)
                            current = molecule_data[mol_key][col_name]
                            # Keep minimum MIC value
                            if current is None or val < current:
                                molecule_data[mol_key][col_name] = val
                        except (ValueError, TypeError):
                            pass

        # Write CSV
        csv_headers = [
            "molecule_id", "name", "smiles", "source_paper", "head_group_class",
            "scaffold_class", "modern_scaffold_tags", "legacy_flags",
            "modern_disinfectant_score", "candidate_tier", "generation_recommendation",
            "MIC_S_aureus", "MIC_E_coli", "MIC_P_aeruginosa", "MIC_C_albicans",
            "HC50", "selectivity_index"
        ]

        with csv_path.open("w", newline="", encoding="utf-8") as f:
            writer = csv.writer(f)
            writer.writerow(csv_headers)

            for mol_key, mol in molecule_lookup.items():
                data = molecule_data[mol_key]

                # Calculate selectivity index
                selectivity = None
                hc50 = data["HC50"]
                if hc50 is not None:
                    mic_values = [
                        data["MIC_S_aureus"],
                        data["MIC_E_coli"],
                        data["MIC_P_aeruginosa"],
                        data["MIC_C_albicans"]
                    ]
                    valid_mics = [v for v in mic_values if v is not None and v > 0]
                    if valid_mics:
                        lowest_mic = min(valid_mics)
                        selectivity = hc50 / lowest_mic

                writer.writerow([
                    mol.get("molecule_id", ""),
                    mol.get("name_as_written", "") or mol.get("normalized_name", ""),
                    mol.get("smiles", ""),
                    mol.get("source_paper", ""),
                    _sanitize_cell(mol.get("head_group_class", "")),
                    _sanitize_cell(mol.get("scaffold_class", "")),
                    _sanitize_cell(mol.get("modern_scaffold_tags", "")),
                    _sanitize_cell(mol.get("legacy_flags", "")),
                    mol.get("modern_disinfectant_score", ""),
                    mol.get("candidate_tier", ""),
                    mol.get("generation_recommendation", ""),
                    data["MIC_S_aureus"] if data["MIC_S_aureus"] is not None else "",
                    data["MIC_E_coli"] if data["MIC_E_coli"] is not None else "",
                    data["MIC_P_aeruginosa"] if data["MIC_P_aeruginosa"] is not None else "",
                    data["MIC_C_albicans"] if data["MIC_C_albicans"] is not None else "",
                    hc50 if hc50 is not None else "",
                    f"{selectivity:.2f}" if selectivity is not None else ""
                ])

        files_created.append(str(csv_path))
    except Exception as e:
        error_msg = f"CSV export failed: {e}"
        errors.append(error_msg)
        if logger:
            logger.log_warning(error_msg, stage="Merge Exports")

    # ========================================
    # 3. Create molecules.smi
    # ========================================
    smi_path = output_dir / "molecules.smi"
    try:
        # Deduplicate by SMILES - keep first occurrence
        seen_smiles = set()
        with smi_path.open("w", encoding="utf-8") as f:
            for mol in all_molecules:
                smiles = mol.get("smiles", "")
                if smiles and smiles not in seen_smiles:
                    seen_smiles.add(smiles)
                    mol_id = mol.get("molecule_id", "unknown")
                    f.write(f"{smiles}\t{mol_id}\n")

        files_created.append(str(smi_path))
    except Exception as e:
        error_msg = f"SMI export failed: {e}"
        errors.append(error_msg)
        if logger:
            logger.log_warning(error_msg, stage="Merge Exports")

    # ========================================
    # 4. Create fairchem_uma_candidates.csv
    # ========================================
    fairchem_path = output_dir / "fairchem_uma_candidates.csv"
    try:
        fairchem_headers = [
            "molecule_id", "name", "smiles", "source_paper", "doc_id",
            "candidate_tier", "modern_scaffold_tags", "legacy_flags",
            "uma_task", "total_charge_hint", "spin_multiplicity_hint",
            "fragment_count", "atom_count", "readiness_status", "readiness_notes"
        ]

        seen_smiles = set()
        with fairchem_path.open("w", newline="", encoding="utf-8") as f:
            writer = csv.writer(f)
            writer.writerow(fairchem_headers)

            for mol in all_molecules:
                smiles = mol.get("smiles", "")
                if not smiles or smiles in seen_smiles:
                    continue
                seen_smiles.add(smiles)

                status, notes = _infer_uma_readiness(mol)
                total_charge = mol.get("formal_charge")
                if total_charge is None:
                    total_charge = ""

                writer.writerow([
                    mol.get("molecule_id", ""),
                    mol.get("name_as_written", "") or mol.get("normalized_name", ""),
                    smiles,
                    mol.get("source_paper", ""),
                    mol.get("doc_id", ""),
                    mol.get("candidate_tier", ""),
                    _sanitize_cell(mol.get("modern_scaffold_tags", "")),
                    _sanitize_cell(mol.get("legacy_flags", "")),
                    "omol",
                    total_charge,
                    1,
                    mol.get("fragment_count", ""),
                    mol.get("atom_count", ""),
                    status,
                    notes,
                ])

        files_created.append(str(fairchem_path))
    except Exception as e:
        error_msg = f"FairChem UMA candidate export failed: {e}"
        errors.append(error_msg)
        if logger:
            logger.log_warning(error_msg, stage="Merge Exports")

    # ========================================
    # 5. Create candidate_generation_seed.smi
    # ========================================
    seed_smi_path = output_dir / "candidate_generation_seed.smi"
    try:
        seen_smiles = set()
        with seed_smi_path.open("w", encoding="utf-8") as f:
            for mol in all_molecules:
                smiles = mol.get("smiles", "")
                if not smiles or smiles in seen_smiles or not is_generation_seed(mol):
                    continue
                seen_smiles.add(smiles)
                mol_id = mol.get("molecule_id", "unknown")
                f.write(f"{smiles}\t{mol_id}\n")

        files_created.append(str(seed_smi_path))
    except Exception as e:
        error_msg = f"Generation seed SMI export failed: {e}"
        errors.append(error_msg)
        if logger:
            logger.log_warning(error_msg, stage="Merge Exports")

    # ========================================
    # 6. Create legacy_or_low_priority_molecules.csv
    # ========================================
    legacy_path = output_dir / "legacy_or_low_priority_molecules.csv"
    try:
        legacy_headers = [
            "molecule_id", "name", "smiles", "source_paper", "doc_id",
            "head_group_class", "modern_disinfectant_score", "candidate_tier",
            "generation_recommendation", "modern_scaffold_tags", "legacy_flags",
            "notes"
        ]

        seen_smiles = set()
        with legacy_path.open("w", newline="", encoding="utf-8") as f:
            writer = csv.writer(f)
            writer.writerow(legacy_headers)

            for mol in all_molecules:
                smiles = mol.get("smiles", "")
                dedupe_key = smiles or f"{mol.get('doc_id', '')}:{mol.get('molecule_id', '')}"
                if dedupe_key in seen_smiles or not is_legacy_or_low_priority(mol):
                    continue
                seen_smiles.add(dedupe_key)

                writer.writerow([
                    mol.get("molecule_id", ""),
                    mol.get("name_as_written", "") or mol.get("normalized_name", ""),
                    smiles,
                    mol.get("source_paper", ""),
                    mol.get("doc_id", ""),
                    _sanitize_cell(mol.get("head_group_class", "")),
                    mol.get("modern_disinfectant_score", ""),
                    mol.get("candidate_tier", ""),
                    mol.get("generation_recommendation", ""),
                    _sanitize_cell(mol.get("modern_scaffold_tags", "")),
                    _sanitize_cell(mol.get("legacy_flags", "")),
                    mol.get("generation_notes", ""),
                ])

        files_created.append(str(legacy_path))
    except Exception as e:
        error_msg = f"Legacy molecule export failed: {e}"
        errors.append(error_msg)
        if logger:
            logger.log_warning(error_msg, stage="Merge Exports")

    # ========================================
    # 7. Create screening_score_profile.json
    # ========================================
    score_profile_path = output_dir / "screening_score_profile.json"
    try:
        score_profile = _build_screening_score_profile(stats)
        score_profile_path.write_text(json.dumps(score_profile, indent=2), encoding="utf-8")
        files_created.append(str(score_profile_path))
    except Exception as e:
        error_msg = f"Screening score profile export failed: {e}"
        errors.append(error_msg)
        if logger:
            logger.log_warning(error_msg, stage="Merge Exports")

    # Log completion
    status = "success" if not errors else "partial"
    if logger:
        logger.log_stage("Merge Exports", {
            "status": status,
            "files_created": files_created,
            "errors": errors,
            "stats": stats
        })
        logger.emit_progress("MERGE", "end", stats=stats)

    return {
        "status": status,
        "files_created": files_created,
        "errors": errors,
        "stats": stats,
        "unified_excel": str(excel_path) if str(excel_path) in files_created else None,
        "analysis_csv": str(csv_path) if str(csv_path) in files_created else None,
        "smiles_file": str(smi_path) if str(smi_path) in files_created else None,
        "fairchem_uma_candidates": str(fairchem_path) if str(fairchem_path) in files_created else None,
        "candidate_generation_seed": str(seed_smi_path) if str(seed_smi_path) in files_created else None,
        "legacy_or_low_priority_molecules": str(legacy_path) if str(legacy_path) in files_created else None,
        "screening_score_profile": str(score_profile_path) if str(score_profile_path) in files_created else None
    }
