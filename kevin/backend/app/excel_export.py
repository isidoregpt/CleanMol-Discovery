"""
Excel export for chemist review.
Creates reviewer-friendly workbooks with molecules, experiments, and results.
"""
import json
from pathlib import Path
from typing import Dict, Any, Optional, List
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter


# Style definitions
HEADER_FONT = Font(bold=True, color="FFFFFF")
HEADER_FILL = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
HEADER_ALIGNMENT = Alignment(horizontal="center", vertical="center", wrap_text=True)

ALT_ROW_FILL = PatternFill(start_color="D9E2F3", end_color="D9E2F3", fill_type="solid")
THIN_BORDER = Border(
    left=Side(style="thin"),
    right=Side(style="thin"),
    top=Side(style="thin"),
    bottom=Side(style="thin")
)

ORPHAN_FILL = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")
NO_SMILES_FILL = PatternFill(start_color="FFEB9C", end_color="FFEB9C", fill_type="solid")


def _auto_width(ws, min_width: int = 10, max_width: int = 50):
    """Auto-fit column widths based on content."""
    for col_idx, column_cells in enumerate(ws.columns, 1):
        max_len = 0
        for cell in column_cells:
            try:
                cell_len = len(str(cell.value or ""))
                if cell_len > max_len:
                    max_len = cell_len
            except Exception:
                pass
        adjusted_width = min(max(max_len + 2, min_width), max_width)
        ws.column_dimensions[get_column_letter(col_idx)].width = adjusted_width


def _apply_header_style(ws, row: int = 1):
    """Apply header styling to first row."""
    for cell in ws[row]:
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = HEADER_ALIGNMENT
        cell.border = THIN_BORDER


def _apply_data_style(ws, start_row: int = 2, highlight_orphans: bool = True):
    """Apply data row styling with alternating colors."""
    for row_idx, row in enumerate(ws.iter_rows(min_row=start_row), start=0):
        for cell in row:
            cell.border = THIN_BORDER
            cell.alignment = Alignment(vertical="center", wrap_text=True)

            # Check for orphan annotation
            if highlight_orphans and cell.value and "_orphan" in str(cell.value).lower():
                cell.fill = ORPHAN_FILL
            elif row_idx % 2 == 1:
                cell.fill = ALT_ROW_FILL


def create_review_workbook(extraction: dict, paper_title: str = "Unknown") -> Workbook:
    """
    Create a reviewer-friendly Excel workbook from extraction data.

    Args:
        extraction: Dict with molecules, experiments, results
        paper_title: Title for the workbook

    Returns:
        openpyxl Workbook object
    """
    wb = Workbook()

    # Remove default sheet
    if "Sheet" in wb.sheetnames:
        del wb["Sheet"]

    # === Molecules Sheet ===
    ws_mol = wb.create_sheet("Molecules")
    mol_headers = [
        "molecule_id", "name_as_written", "normalized_name", "smiles",
        "smiles_source", "role", "notes", "orphan_status"
    ]
    ws_mol.append(mol_headers)

    molecules = extraction.get("molecules") or []
    for mol in molecules:
        meta = mol.get("_meta") or {}
        orphan_status = ""
        if meta.get("orphan_molecule"):
            orphan_status = f"ORPHAN (orig: {meta.get('original_molecule_id', '?')})"

        ws_mol.append([
            mol.get("molecule_id", ""),
            mol.get("name_as_written", ""),
            mol.get("normalized_name", ""),
            mol.get("smiles", ""),
            mol.get("smiles_source", ""),
            mol.get("role", ""),
            mol.get("notes", ""),
            orphan_status
        ])

    _apply_header_style(ws_mol)
    _apply_data_style(ws_mol)
    _auto_width(ws_mol)

    # Highlight molecules without SMILES
    for row_idx, row in enumerate(ws_mol.iter_rows(min_row=2), start=2):
        smiles_cell = row[3]  # smiles column
        if not smiles_cell.value:
            for cell in row:
                if not cell.fill or cell.fill.start_color.index == "00000000":
                    cell.fill = NO_SMILES_FILL

    # === Experiments Sheet ===
    ws_exp = wb.create_sheet("Experiments")
    exp_headers = [
        "experiment_id", "type", "conditions", "equipment",
        "parameters", "notes", "orphan_status"
    ]
    ws_exp.append(exp_headers)

    experiments = extraction.get("experiments") or []
    for exp in experiments:
        meta = exp.get("_meta") or {}
        orphan_status = ""
        if meta.get("orphan_experiment"):
            orphan_status = f"ORPHAN (orig: {meta.get('original_experiment_id', '?')})"

        # Serialize complex fields
        conditions = exp.get("conditions", "")
        if isinstance(conditions, dict):
            conditions = json.dumps(conditions, ensure_ascii=False)

        parameters = exp.get("parameters", "")
        if isinstance(parameters, dict):
            parameters = json.dumps(parameters, ensure_ascii=False)

        ws_exp.append([
            exp.get("experiment_id", ""),
            exp.get("type", ""),
            conditions,
            exp.get("equipment", ""),
            parameters,
            exp.get("notes", ""),
            orphan_status
        ])

    _apply_header_style(ws_exp)
    _apply_data_style(ws_exp)
    _auto_width(ws_exp)

    # === Results Sheet ===
    ws_res = wb.create_sheet("Results")
    res_headers = [
        "result_id", "molecule_id", "experiment_id", "property_name",
        "value", "unit", "conditions", "notes", "orphan_status"
    ]
    ws_res.append(res_headers)

    results = extraction.get("results") or []
    for res in results:
        meta = res.get("_meta") or {}
        orphan_parts = []
        if meta.get("orphan_experiment"):
            orphan_parts.append(f"exp_orphan (orig: {meta.get('original_experiment_id', '?')})")
        if meta.get("orphan_molecule"):
            orphan_parts.append(f"mol_orphan (orig: {meta.get('original_molecule_id', '?')})")
        orphan_status = "; ".join(orphan_parts)

        # Serialize complex fields
        conditions = res.get("conditions", "")
        if isinstance(conditions, dict):
            conditions = json.dumps(conditions, ensure_ascii=False)

        ws_res.append([
            res.get("result_id", ""),
            res.get("molecule_id", ""),
            res.get("experiment_id", ""),
            res.get("property_name", ""),
            res.get("value", ""),
            res.get("unit", ""),
            conditions,
            res.get("notes", ""),
            orphan_status
        ])

    _apply_header_style(ws_res)
    _apply_data_style(ws_res)
    _auto_width(ws_res)

    # === Summary Sheet (first position) ===
    ws_summary = wb.create_sheet("Summary", 0)
    ws_summary.append(["Paper Title", paper_title])
    ws_summary.append([])
    ws_summary.append(["Entity", "Count"])
    ws_summary.append(["Molecules", len(molecules)])
    ws_summary.append(["Experiments", len(experiments)])
    ws_summary.append(["Results", len(results)])
    ws_summary.append([])

    # SMILES coverage
    with_smiles = sum(1 for m in molecules if m.get("smiles"))
    ws_summary.append(["SMILES Coverage", f"{with_smiles}/{len(molecules)}"])

    # Orphan counts
    orphan_results = sum(1 for r in results if (r.get("_meta") or {}).get("orphan_experiment") or (r.get("_meta") or {}).get("orphan_molecule"))
    ws_summary.append(["Orphan Results", orphan_results])

    _apply_header_style(ws_summary, row=1)
    _apply_header_style(ws_summary, row=3)
    _auto_width(ws_summary)

    return wb


def export_combined_workbook(
    extractions: List[Dict[str, Any]],
    output_path: Path,
    logger=None
) -> Path:
    """
    Export a combined workbook from multiple paper extractions.

    Args:
        extractions: List of dicts with 'extraction' and 'paper_title' keys
        output_path: Path to save the workbook
        logger: Optional PipelineLogger

    Returns:
        Path to the saved workbook
    """
    wb = Workbook()

    # Remove default sheet
    if "Sheet" in wb.sheetnames:
        del wb["Sheet"]

    # Combine all data
    all_molecules = []
    all_experiments = []
    all_results = []
    paper_summaries = []

    for item in extractions:
        extraction = item.get("extraction") or item
        paper_title = item.get("paper_title", "Unknown")

        molecules = extraction.get("molecules") or []
        experiments = extraction.get("experiments") or []
        results = extraction.get("results") or []

        # Add source paper to each record
        for mol in molecules:
            mol["source_paper"] = paper_title
            all_molecules.append(mol)

        for exp in experiments:
            exp["source_paper"] = paper_title
            all_experiments.append(exp)

        for res in results:
            res["source_paper"] = paper_title
            all_results.append(res)

        with_smiles = sum(1 for m in molecules if m.get("smiles"))
        paper_summaries.append({
            "title": paper_title,
            "molecules": len(molecules),
            "experiments": len(experiments),
            "results": len(results),
            "smiles_coverage": f"{with_smiles}/{len(molecules)}"
        })

    # === Summary Sheet ===
    ws_summary = wb.create_sheet("Summary", 0)
    ws_summary.append(["Combined Extraction Summary"])
    ws_summary.append([])
    ws_summary.append(["Total Molecules", len(all_molecules)])
    ws_summary.append(["Total Experiments", len(all_experiments)])
    ws_summary.append(["Total Results", len(all_results)])
    ws_summary.append(["Papers Processed", len(extractions)])
    ws_summary.append([])
    ws_summary.append(["Paper", "Molecules", "Experiments", "Results", "SMILES Coverage"])

    for ps in paper_summaries:
        ws_summary.append([
            ps["title"], ps["molecules"], ps["experiments"],
            ps["results"], ps["smiles_coverage"]
        ])

    _apply_header_style(ws_summary, row=1)
    _apply_header_style(ws_summary, row=8)
    _auto_width(ws_summary)

    # === Molecules Sheet ===
    ws_mol = wb.create_sheet("Molecules")
    mol_headers = [
        "source_paper", "molecule_id", "name_as_written", "normalized_name",
        "smiles", "smiles_source", "role", "notes"
    ]
    ws_mol.append(mol_headers)

    for mol in all_molecules:
        ws_mol.append([
            mol.get("source_paper", ""),
            mol.get("molecule_id", ""),
            mol.get("name_as_written", ""),
            mol.get("normalized_name", ""),
            mol.get("smiles", ""),
            mol.get("smiles_source", ""),
            mol.get("role", ""),
            mol.get("notes", "")
        ])

    _apply_header_style(ws_mol)
    _apply_data_style(ws_mol, highlight_orphans=False)
    _auto_width(ws_mol)

    # Highlight molecules without SMILES
    for row_idx, row in enumerate(ws_mol.iter_rows(min_row=2), start=2):
        smiles_cell = row[4]  # smiles column (0-indexed in row)
        if not smiles_cell.value:
            for cell in row:
                if not cell.fill or cell.fill.start_color.index == "00000000":
                    cell.fill = NO_SMILES_FILL

    # === Experiments Sheet ===
    ws_exp = wb.create_sheet("Experiments")
    exp_headers = [
        "source_paper", "experiment_id", "type", "conditions",
        "equipment", "parameters", "notes"
    ]
    ws_exp.append(exp_headers)

    for exp in all_experiments:
        conditions = exp.get("conditions", "")
        if isinstance(conditions, dict):
            conditions = json.dumps(conditions, ensure_ascii=False)

        parameters = exp.get("parameters", "")
        if isinstance(parameters, dict):
            parameters = json.dumps(parameters, ensure_ascii=False)

        ws_exp.append([
            exp.get("source_paper", ""),
            exp.get("experiment_id", ""),
            exp.get("type", ""),
            conditions,
            exp.get("equipment", ""),
            parameters,
            exp.get("notes", "")
        ])

    _apply_header_style(ws_exp)
    _apply_data_style(ws_exp, highlight_orphans=False)
    _auto_width(ws_exp)

    # === Results Sheet ===
    ws_res = wb.create_sheet("Results")
    res_headers = [
        "source_paper", "result_id", "molecule_id", "experiment_id",
        "property_name", "value", "unit", "conditions", "notes"
    ]
    ws_res.append(res_headers)

    for res in all_results:
        conditions = res.get("conditions", "")
        if isinstance(conditions, dict):
            conditions = json.dumps(conditions, ensure_ascii=False)

        ws_res.append([
            res.get("source_paper", ""),
            res.get("result_id", ""),
            res.get("molecule_id", ""),
            res.get("experiment_id", ""),
            res.get("property_name", ""),
            res.get("value", ""),
            res.get("unit", ""),
            conditions,
            res.get("notes", "")
        ])

    _apply_header_style(ws_res)
    _apply_data_style(ws_res, highlight_orphans=False)
    _auto_width(ws_res)

    # Save workbook
    output_path = Path(output_path)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output_path)

    if logger:
        logger.log_stage("Excel Export", {"status": "success", "output": str(output_path)})

    return output_path
